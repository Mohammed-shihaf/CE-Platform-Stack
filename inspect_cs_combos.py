import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\moham\Downloads\Testable_Combo_Details_Filled (1).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

sheets = ['2-Lang Combinations', '3-Lang Combinations', '4-Lang Combinations', '5-Lang Combination']

for s in sheets:
    if s not in wb.sheetnames:
        continue
    ws = wb[s]
    print(f"\n=================== SHEET: {s} ===================")
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        row_str = str(row)
        if 'C#' in row_str or 'CS' in row_str:
            print(f"Row {r_idx}: {row[:10]}")
