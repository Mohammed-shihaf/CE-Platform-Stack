import openpyxl

file_path = r"C:\Users\moham\Downloads\Testable_Combo_Details_Filled (1).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

print("Sheet names:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    print(f"\n=================== SHEET: {sheet_name} ===================")
    ws = wb[sheet_name]
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if r_idx <= 40: # print first 40 rows per sheet
            print(f"Row {r_idx}: {row}")
