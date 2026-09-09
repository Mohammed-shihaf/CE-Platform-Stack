import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\moham\Downloads\Testable_Combo_Details_Filled (1).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

sheets_to_check = ['Tool Matrix by Language', 'Metric Derivation & Proof', '2-Lang Combinations', 'Combo Master Index']

for sheet_name in sheets_to_check:
    if sheet_name not in wb.sheetnames:
        continue
    print(f"\n=================== SHEET: {sheet_name} ===================")
    ws = wb[sheet_name]
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if r_idx <= 60:
            print(f"Row {r_idx}: {row}")
