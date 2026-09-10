import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\moham\Downloads\Testable_Combo_Details_Filled (1).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

ws = wb['Tool Matrix by Language']
headers = [cell.value for cell in ws[1]]
print("Headers:", headers)

print("\nC# Tools in Sheet:")
for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if r_idx == 1:
        continue
    # Check if row is about C#
    row_vals = [str(v) for v in row if v is not None]
    if any('c#' in v.lower() or 'csharp' in v.lower() or 'dotnet' in v.lower() for v in row_vals):
        print(f"Row {r_idx}: {row}")
