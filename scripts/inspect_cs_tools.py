import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

file_path = r"C:\Users\moham\Downloads\Testable_Combo_Details_Filled (1).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

print("=== TOOL MATRIX BY LANGUAGE ===")
ws = wb['Tool Matrix by Language']
for row in ws.iter_rows(values_only=True):
    # Print rows relevant to C#
    row_str = str(row)
    if 'C#' in row_str or 'csharp' in row_str.lower() or 'dotnet' in row_str.lower() or 'roslyn' in row_str.lower() or 'sonarqube' in row_str.lower() or 'security' in row_str.lower():
        print(row)

print("\n=== METRIC DERIVATION & PROOF (C# entries) ===")
ws_metrics = wb['Metric Derivation & Proof']
for row_idx, row in enumerate(ws_metrics.iter_rows(values_only=True), start=1):
    row_str = str(row)
    if 'C#' in row_str or 'csharp' in row_str.lower() or 'dotnet' in row_str.lower():
        print(f"Row {row_idx}: {row[:6]}")
